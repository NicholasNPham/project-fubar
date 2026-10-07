"""Tests for excel_io.py, on small generated workbooks. No real data."""

import hashlib
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from checker import EXCLUDE_GOOD
from excel_io import (
    conflicting_image_ids,
    group_by_case,
    output_path_for,
    resume_output,
    start_output,
)
from exceptions import InputProblem
from models import CheckResult, ImageRow, Status

NOW = datetime(2026, 10, 6, 12, 0, 0)

# Sample rows from CLAUDE.md. 15072070 is an int, 15031454 a float the way
# Excel often returns it, and 15076194 text, so normalization is exercised.
SAMPLE = [
    ("CF12010742XX", 15072070, None),
    ("CF11003819XX", 15031454.0, None),
    ("CF1400494AXXS", "15076194", None),
    ("CF1400494AXXS", 15080911, None),
]


def _make_input(path: Path, headers: tuple, rows: list) -> Path:
    """Write a one-sheet workbook with the given header and rows."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    return path


def _digest(path: Path) -> str:
    """Hash a file so any change to it shows up."""
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture
def input_file(tmp_path: Path) -> Path:
    """The CLAUDE.md sample rows with the real sheet's headers."""
    return _make_input(tmp_path / "input.xlsx", ("Case Number", "Image Id", "Exclude"), SAMPLE)


def _result(status: str, exclude: str = "") -> CheckResult:
    """A CheckResult with fixed filler values."""
    return CheckResult(status, exclude, "16.72 MB" if exclude else "", "Local", NOW)


def test_output_name() -> None:
    """Output is <name>_checked_<YYYYMMDD_HHMMSS>.xlsx next to the input."""
    assert output_path_for(Path("C:/x/list.xlsx"), NOW) == Path("C:/x/list_checked_20261006_120000.xlsx")


def test_input_file_is_never_changed(input_file: Path) -> None:
    """Writing and saving results touches only the copy."""
    before = _digest(input_file)
    book = start_output(input_file, NOW)
    rows, _ = book.read_rows()
    book.write_result(rows[0].row_number, _result(Status.OK, EXCLUDE_GOOD))
    book.save()
    assert _digest(input_file) == before
    assert book.path != input_file


def test_reads_and_normalizes_rows(input_file: Path) -> None:
    """Int, float and text Image Ids all come out as plain digits."""
    rows, bad = start_output(input_file, NOW).read_rows()
    assert bad == {}
    assert [(r.row_number, r.case_number, r.image_id) for r in rows] == [
        (2, "CF12010742XX", "15072070"),
        (3, "CF11003819XX", "15031454"),
        (4, "CF1400494AXXS", "15076194"),
        (5, "CF1400494AXXS", "15080911"),
    ]


def test_headers_matched_by_name_not_position(tmp_path: Path) -> None:
    """Reordered, padded, odd-case headers still match."""
    path = _make_input(
        tmp_path / "input.xlsx",
        ("exclude", "  IMAGE ID ", "Notes", "case number"),
        [(None, 15072070, "x", "CF12010742XX")],
    )
    rows, _ = start_output(path, NOW).read_rows()
    assert rows == [ImageRow(2, "CF12010742XX", "15072070")]


def test_missing_column_stops_the_run(tmp_path: Path) -> None:
    """No Exclude column is an input problem, caught before any browser work."""
    path = _make_input(tmp_path / "input.xlsx", ("Case Number", "Image Id"), [("CF1", 1)])
    with pytest.raises(InputProblem):
        start_output(path, NOW)


def test_bad_rows_fail_alone(tmp_path: Path) -> None:
    """A bad cell becomes that row's ERROR; blank rows are skipped silently."""
    path = _make_input(
        tmp_path / "input.xlsx",
        ("Case Number", "Image Id", "Exclude"),
        [("CF1", "abc", None), (None, 15072070, None), (None, None, None), ("CF2", 15031454, None)],
    )
    rows, bad = start_output(path, NOW).read_rows()
    assert [r.row_number for r in rows] == [5]
    assert set(bad) == {2, 3}
    assert all(status.startswith(Status.ERROR_PREFIX) for status in bad.values())


def test_results_written_and_resume_skips_finished_rows(input_file: Path) -> None:
    """OK and CORRUPTED rows are done on resume; ERROR and unchecked rows are not."""
    book = start_output(input_file, NOW)
    book.write_result(2, _result(Status.OK, EXCLUDE_GOOD))
    book.write_result(3, _result(Status.CORRUPTED))
    book.write_result(4, _result(Status.error("details did not load for this image")))
    book.save()

    sheet = load_workbook(book.path).worksheets[0]
    assert [c.value for c in sheet[1]] == [
        "Case Number", "Image Id", "Exclude", "Status", "File Size", "Storage Location", "Checked At",
    ]
    assert sheet["C2"].value == EXCLUDE_GOOD
    assert sheet["C3"].value is None

    resumed = resume_output(input_file, book.path)
    assert [resumed.is_done(n) for n in (2, 3, 4, 5)] == [True, True, False, False]
    # Resuming must reuse the output columns, not add a second set.
    assert load_workbook(book.path).worksheets[0].max_column == 7


def test_resume_refuses_the_input_file(input_file: Path) -> None:
    """Passing the input as the resume file would write over it."""
    with pytest.raises(InputProblem):
        resume_output(input_file, input_file)


def test_group_by_case_keeps_order_and_merges_repeats() -> None:
    """One entry per case, one per Image Id, repeated rows grouped together."""
    rows = [
        ImageRow(2, "B", "1"),
        ImageRow(3, "A", "2"),
        ImageRow(4, "B", "3"),
        ImageRow(5, "B", "1"),
    ]
    grouped = group_by_case(rows)
    assert list(grouped) == ["B", "A"]
    assert list(grouped["B"]) == ["1", "3"]
    assert [r.row_number for r in grouped["B"]["1"]] == [2, 5]


def test_conflicting_image_ids() -> None:
    """An Image Id under two different cases is flagged; a repeat on one case is not."""
    rows = [ImageRow(2, "A", "1"), ImageRow(3, "B", "1"), ImageRow(4, "C", "2"), ImageRow(5, "C", "2")]
    assert conflicting_image_ids(rows) == {"1"}
