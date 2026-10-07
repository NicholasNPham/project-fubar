"""Reads the input sheet and writes results to a separate checked copy.

The input file is never written. A run opens it, saves it straight away
under <name>_checked_<stamp>.xlsx, and only ever touches that copy. A resumed
run opens the copy instead and skips rows it already finished.
"""

import os
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from checker import normalize_image_id
from exceptions import InputProblem, SystemProblem
from models import CheckResult, ImageRow, Status

# Input headers, matched case-insensitive and trimmed, never by column letter.
CASE_HEADER = "Case Number"
IMAGE_HEADER = "Image Id"
EXCLUDE_HEADER = "Exclude"
REQUIRED_HEADERS = (CASE_HEADER, IMAGE_HEADER, EXCLUDE_HEADER)

# Added after the original columns, or reused if a resumed copy has them.
STATUS_HEADER = "Status"
FILE_SIZE_HEADER = "File Size"
STORAGE_HEADER = "Storage Location"
CHECKED_AT_HEADER = "Checked At"
OUTPUT_HEADERS = (STATUS_HEADER, FILE_SIZE_HEADER, STORAGE_HEADER, CHECKED_AT_HEADER)

HEADER_ROW = 1
OUTPUT_STAMP_FORMAT = "%Y%m%d_%H%M%S"


def output_path_for(input_path: Path, now: datetime) -> Path:
    """Name the checked copy: <original name>_checked_<YYYYMMDD_HHMMSS>.xlsx.

    Args:
        input_path: The input spreadsheet.
        now: Run start time, for the stamp.

    Returns:
        A path next to the input file.
    """
    return input_path.with_name(f"{input_path.stem}_checked_{now:{OUTPUT_STAMP_FORMAT}}.xlsx")


def _key(header: object) -> str:
    """Normalize a header cell for matching: trimmed, lower case."""
    return str(header).strip().lower() if header is not None else ""


class OutputBook:
    """The checked copy: where rows are read from and results written to.

    Build one with start_output or resume_output, never directly, so the
    input-file guard always runs.
    """

    def __init__(self, path: Path, input_path: Path) -> None:
        """Open a workbook that is already the output copy, and find its columns.

        Args:
            path: The output copy to read and write.
            input_path: The original input, only to refuse writing over it.

        Raises:
            InputProblem: If path is the input file, cannot be opened, or is
                missing a required column.
        """
        if path.resolve() == input_path.resolve():
            raise InputProblem(f"Refusing to write to the input file itself: {path}")

        try:
            self._workbook = load_workbook(path)
        except (OSError, ValueError, KeyError) as error:
            raise InputProblem(f"Could not open {path}: {error}") from error

        self.path = path
        # The sheet is a single Sheet1. First sheet, not "active", because
        # active is whichever tab was selected when the file was last saved.
        self._sheet: Worksheet = self._workbook.worksheets[0]
        self._columns = self._find_columns()

    def _find_columns(self) -> dict[str, int]:
        """Map each header this tool uses to its column number, adding output ones.

        Returns:
            Lower-case header to 1-based column number.

        Raises:
            InputProblem: If Case Number, Image Id or Exclude is missing.
        """
        columns: dict[str, int] = {}
        for cell in self._sheet[HEADER_ROW]:
            key = _key(cell.value)
            if key and key not in columns:
                columns[key] = cell.column

        missing = [h for h in REQUIRED_HEADERS if _key(h) not in columns]
        if missing:
            raise InputProblem(f"Input sheet has no column named: {', '.join(missing)}")

        for header in OUTPUT_HEADERS:
            if _key(header) not in columns:
                column = self._sheet.max_column + 1
                self._sheet.cell(row=HEADER_ROW, column=column, value=header)
                columns[_key(header)] = column
        return columns

    def _value(self, row_number: int, header: str) -> object:
        """Read one cell by header name."""
        return self._sheet.cell(row=row_number, column=self._columns[_key(header)]).value

    def read_rows(self) -> tuple[list[ImageRow], dict[int, str]]:
        """Read every data row that has a case number or an Image Id.

        A bad cell fails only its own row, not the run: 2,000 good rows should
        not wait on one typo.

        Returns:
            (rows to check, {row_number: ERROR status} for rows that cannot be).
            Fully blank rows are in neither.
        """
        rows: list[ImageRow] = []
        bad: dict[int, str] = {}

        for row_number in range(HEADER_ROW + 1, self._sheet.max_row + 1):
            case_raw = self._value(row_number, CASE_HEADER)
            image_raw = self._value(row_number, IMAGE_HEADER)
            case_number = str(case_raw).strip() if case_raw is not None else ""

            if not case_number and (image_raw is None or not str(image_raw).strip()):
                continue
            if not case_number:
                bad[row_number] = Status.error("no case number")
                continue

            try:
                image_id = normalize_image_id(image_raw)
            except InputProblem as error:
                bad[row_number] = Status.error(str(error))
                continue

            rows.append(ImageRow(row_number, case_number, image_id))

        return rows, bad

    def is_done(self, row_number: int) -> bool:
        """Whether a resumed run can skip this row.

        Any Status except an ERROR counts as done. ERROR rows are retried,
        since most of them are STAC hiccups that a second pass gets through.

        Args:
            row_number: Sheet row.

        Returns:
            True if the row already has a non-ERROR Status.
        """
        status = self._value(row_number, STATUS_HEADER)
        status = str(status).strip() if status is not None else ""
        return bool(status) and not status.startswith(Status.ERROR_PREFIX)

    def write_result(self, row_number: int, result: CheckResult) -> None:
        """Write one verdict to its row. Overwrites any Exclude already there.

        Args:
            row_number: Sheet row.
            result: The verdict.
        """
        values = {
            EXCLUDE_HEADER: result.exclude or None,
            STATUS_HEADER: result.status,
            FILE_SIZE_HEADER: result.file_size or None,
            STORAGE_HEADER: result.storage_location or None,
            CHECKED_AT_HEADER: result.checked_at,
        }
        for header, value in values.items():
            self._sheet.cell(row=row_number, column=self._columns[_key(header)], value=value)

    def save(self) -> None:
        """Save the copy without ever leaving a half-written file behind.

        Writes to a temp file and swaps it in, so a crash or a full disk
        mid-save leaves the previous checkpoint intact.

        Raises:
            SystemProblem: If the save fails, most often because the copy is
                open in Excel. The caller decides whether to carry on.
        """
        temp = self.path.with_name(f"{self.path.stem}.tmp.xlsx")
        try:
            self._workbook.save(temp)
        except OSError as error:
            raise SystemProblem(f"Could not write {temp.name}: {error}") from error
        try:
            os.replace(temp, self.path)
        except OSError as error:
            # The temp file is complete at this point, so nothing is lost.
            # Saying where it is matters most at the end of a long run.
            raise SystemProblem(
                f"Could not update {self.path.name}, probably because it is open in Excel. "
                f"The latest results are in {temp.name}. Close Excel and the next save replaces it."
            ) from error


def start_output(input_path: Path, now: datetime) -> OutputBook:
    """Copy the input to a new checked file and open that copy.

    Args:
        input_path: The input spreadsheet. Read only.
        now: Run start time, for the output name.

    Returns:
        The opened copy.

    Raises:
        InputProblem: If the input cannot be read, lacks a required column,
            or the output name is already taken.
    """
    output_path = output_path_for(input_path, now)
    if output_path.exists():
        raise InputProblem(f"Output file already exists: {output_path}")

    try:
        load_workbook(input_path).save(output_path)
    except (OSError, ValueError, KeyError) as error:
        raise InputProblem(f"Could not copy {input_path} to {output_path}: {error}") from error

    return OutputBook(output_path, input_path)


def resume_output(input_path: Path, output_path: Path) -> OutputBook:
    """Reopen a partly finished checked copy.

    Args:
        input_path: The original input, only to refuse resuming onto it.
        output_path: The checked copy from an earlier run.

    Returns:
        The opened copy.

    Raises:
        InputProblem: If output_path is the input file, does not exist, or
            lacks a required column.
    """
    if not output_path.exists():
        raise InputProblem(f"Nothing to resume, no file at {output_path}")
    return OutputBook(output_path, input_path)


def group_by_case(rows: list[ImageRow]) -> dict[str, dict[str, list[ImageRow]]]:
    """Group rows by case number, then by Image Id, keeping first-seen order.

    So each case is searched once and each image on it read once. The inner
    list holds every row with that Image Id on that case, which is normally
    one; a repeated row just gets the same result.

    Args:
        rows: Rows to check.

    Returns:
        {case_number: {image_id: [ImageRow, ...]}}.
    """
    grouped: dict[str, dict[str, list[ImageRow]]] = {}
    for row in rows:
        grouped.setdefault(row.case_number, {}).setdefault(row.image_id, []).append(row)
    return grouped


def conflicting_image_ids(rows: list[ImageRow]) -> set[str]:
    """Image Ids listed under more than one case number.

    An Image Id belongs to exactly one case, so this is a data problem in the
    sheet, and checking either case would be a guess about which one is right.

    Args:
        rows: Rows to check.

    Returns:
        The Image Ids that appear with two or more different case numbers.
    """
    cases_for: dict[str, set[str]] = {}
    for row in rows:
        cases_for.setdefault(row.image_id, set()).add(row.case_number)
    return {image_id for image_id, cases in cases_for.items() if len(cases) > 1}
